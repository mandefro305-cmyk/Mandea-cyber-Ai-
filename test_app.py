import io
import os
import unittest
from PIL import Image
import docx

from file_parser import (
    parse_text_file,
    parse_pdf_file,
    parse_docx_file,
    encode_image_to_base64,
    process_uploaded_file
)
from api_client import prepare_messages_for_api, fetch_available_models
from config import estimate_tokens, calculate_cost
from security_utils import redact_sensitive_data, scan_code_for_vulnerabilities
from data_utils import parse_csv_file
from session_utils import export_chat_to_markdown, export_chat_to_json
from db_utils import init_db, create_session, get_all_sessions, save_message, get_session_messages, delete_session
from rag_utils import chunk_text, search_chunks
from remediation_utils import generate_remediation_diff
from search_utils import format_search_context

class DummyUploadedFile:
    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data

    def getvalue(self) -> bytes:
        return self._data

class TestAppAllFeatures(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        init_db()

    def test_parse_text_file(self):
        content = b"Hello, world! This is a test file."
        res = parse_text_file(content, "test.txt")
        self.assertEqual(res, "Hello, world! This is a test file.")

    def test_parse_docx_file(self):
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1")
        doc.add_paragraph("Paragraph 2")
        bio = io.BytesIO()
        doc.save(bio)
        res = parse_docx_file(bio.getvalue(), "test.docx")
        self.assertIn("Paragraph 1", res)
        self.assertIn("Paragraph 2", res)

    def test_encode_image_to_base64(self):
        img = Image.new("RGB", (10, 10), color="red")
        bio = io.BytesIO()
        img.save(bio, format="PNG")
        b64_str, mime_type = encode_image_to_base64(bio.getvalue())
        self.assertEqual(mime_type, "image/png")
        self.assertTrue(len(b64_str) > 0)

    def test_process_uploaded_file_doc(self):
        dummy = DummyUploadedFile("notes.txt", b"Test notes content")
        res = process_uploaded_file(dummy)
        self.assertEqual(res["filename"], "notes.txt")
        self.assertEqual(res["file_type"], "document")
        self.assertEqual(res["content"], "Test notes content")

    def test_prepare_messages_for_api(self):
        attached = [
            {
                "filename": "doc.txt",
                "file_type": "document",
                "content": "Doc content"
            }
        ]
        history = [
            {"role": "user", "content": "Prior message"},
            {"role": "assistant", "content": "Prior response"}
        ]
        user_prompt = "Summarize this document"
        messages = prepare_messages_for_api(history, attached, user_prompt)

        self.assertEqual(len(messages), 3)
        self.assertEqual(messages[0]["role"], "user")
        self.assertEqual(messages[1]["role"], "assistant")
        self.assertEqual(messages[2]["role"], "user")

    def test_fetch_available_models_fallback(self):
        models = fetch_available_models(api_key="", base_url="https://agentrouter.ai/v1")
        self.assertTrue(len(models) > 0)
        self.assertIn("gpt-4o", models)

    def test_redact_sensitive_data(self):
        sample = "My api_key='sk-12345678901234567890' and email test@example.com"
        redacted, count = redact_sensitive_data(sample)
        self.assertTrue(count >= 2)
        self.assertNotIn("sk-12345678901234567890", redacted)

    def test_scan_code_for_vulnerabilities(self):
        vulnerable_code = "eval('import os');\nquery = 'SELECT * FROM users WHERE id=' + user_id"
        findings = scan_code_for_vulnerabilities(vulnerable_code)
        self.assertTrue(len(findings) >= 2)

    def test_db_persistence(self):
        test_sid = "test_s1"
        create_session(test_sid, "Test Session")
        save_message(test_sid, "user", "Hello DB")
        save_message(test_sid, "assistant", "Hello User")

        msgs = get_session_messages(test_sid)
        self.assertEqual(len(msgs), 2)
        self.assertEqual(msgs[0]["content"], "Hello DB")

        delete_session(test_sid)
        self.assertEqual(len(get_session_messages(test_sid)), 0)

    def test_rag_chunking_and_search(self):
        text = "Python is a programming language. Security auditing ensures safety. SQL injection is a common flaw."
        chunks = chunk_text(text, chunk_size=5, overlap=1)
        doc_chunks = [{"source": "test.txt", "text": c} for c in chunks]

        results = search_chunks("SQL injection safety", doc_chunks, top_k=2)
        self.assertTrue(len(results) > 0)
        self.assertIn("source", results[0])

    def test_remediation_diff(self):
        code = "eval('bad_code')"
        findings = ["Critical eval()"]
        diff = generate_remediation_diff("vulnerable.py", code, findings)
        self.assertIn("--- a/vulnerable.py", diff)
        self.assertIn("FIX: Replaced unsafe eval()", diff)

    def test_format_search_context(self):
        search_results = [{"title": "Test Title", "href": "https://example.com", "body": "Test Snippet"}]
        ctx = format_search_context(search_results)
        self.assertIn("--- Live Web Search Context ---", ctx)
        self.assertIn("https://example.com", ctx)

if __name__ == "__main__":
    unittest.main()
